#!/usr/bin/env python3
"""Inventory every sheet, sprint task block, capacity block, and workbook control."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit("audit_workbook.py requires openpyxl: python3 -m pip install openpyxl") from exc


HEADER_ALIASES = {
    "epic": {"эпик", "релиз", "задача релиза", "задачи по релизу", "название доработки"},
    "task": {"название задачи", "наименование задачи"},
    "subtask": {"название позадачи (если есть)", "название подзадачи (если есть)", "подзадача"},
    "description": {"описание задачи", "описание", "документ"},
    "estimate": {"оценка (час)", "оценка", "оценка, ч", "часы"},
    "actual_estimate": {"реальная оценка", "фактическая оценка", "факт (час)", "факт, ч"},
    "direction": {"направление", "тип", "стек"},
    "executor": {"исполнитель", "разработчик", "executor"},
    "priority": {"приоритет", "приоритет внутри спринта"},
    "commitment": {"важность реализации в рамках спринта", "обязательность"},
    "level": {"уровень"},
    "task_id": {"номер задачи", "jira", "tracker", "id"},
    "status": {"статус на проде", "статус", "прод"},
    "section": {"раздел"},
}

SUMMARY_MARKERS = {"разработка", "итого", "итого разработка", "всего", "команда разработки"}
FORMULA_ERROR_PATTERN = re.compile(r"#(?:REF!|ERROR!|VALUE!|NAME\?|N/A|DIV/0!|NUM!|NULL!)", re.IGNORECASE)
TASK_ID_PATTERN = re.compile(r"\b[A-ZА-Я][A-ZА-Я0-9]+-\d+\b", re.IGNORECASE)
URL_PATTERN = re.compile(r"https?://[^\s<>\"]+")
SUMIF_PATTERN = re.compile(
    r"SUMIF\s*\(\s*([^,;]+)\s*[,;]\s*(?:\"(?:[^\"]|\"\")*\"|[^,;]+)\s*[,;]\s*([^)]+)\)",
    re.IGNORECASE,
)


def normalize(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower().replace("ё", "е"))


NORMALIZED_ALIASES = {
    field: {normalize(alias) for alias in aliases}
    for field, aliases in HEADER_ALIASES.items()
}


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def clean_url(value: str) -> str:
    url = value.rstrip(".,;:!?")
    pairs = (("(", ")"), ("[", "]"), ("{", "}"))
    changed = True
    while changed and url:
        changed = False
        for opening, closing in pairs:
            if url.endswith(closing) and url.count(closing) > url.count(opening):
                url = url[:-1].rstrip(".,;:!?")
                changed = True
    return url


def sheet_fingerprint(sheet: Any, last_row: int, last_column: int) -> str:
    cells = []
    for row in sheet.iter_rows(max_row=last_row or 1, max_col=last_column or 1):
        for cell in row:
            if cell.value in (None, "") and not cell.hyperlink and not cell.comment:
                continue
            cells.append({
                "cell": cell.coordinate,
                "value": json_value(cell.value),
                "hyperlink": (
                    {"target": cell.hyperlink.target, "location": cell.hyperlink.location}
                    if cell.hyperlink else None
                ),
                "comment": (
                    cell.comment.text
                    if cell.comment else None
                ),
            })
    payload = {
        "state": sheet.sheet_state,
        "cells": cells,
        "hidden_rows": [row for row in range(1, last_row + 1) if sheet.row_dimensions[row].hidden],
        "hidden_columns": hidden_columns(sheet, last_column),
        "validations": validations(sheet),
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()


def meaningful_bounds(sheet: Any) -> tuple[int, int, int]:
    last_row = 0
    last_column = 0
    nonempty_rows = 0
    for row in sheet.iter_rows():
        used = [cell.column for cell in row if cell.value not in (None, "") or cell.hyperlink or cell.comment]
        if used:
            nonempty_rows += 1
            last_row = row[0].row
            last_column = max(last_column, max(used))
    return last_row, last_column, nonempty_rows


def header_mapping(sheet: Any, row_number: int, last_column: int) -> dict[str, int] | None:
    values = {
        column: normalize(sheet.cell(row_number, column).value)
        for column in range(1, last_column + 1)
        if sheet.cell(row_number, column).value not in (None, "")
    }
    mapping: dict[str, int] = {}
    for field, aliases in NORMALIZED_ALIASES.items():
        for column, value in values.items():
            if value in aliases:
                mapping.setdefault(field, column)

    ambiguous_tasks = [column for column, value in values.items() if value == "задача"]
    if "название доработки" in values.values():
        if ambiguous_tasks:
            mapping["task"] = ambiguous_tasks[0]
            mapping.setdefault("description", ambiguous_tasks[0])
        if len(ambiguous_tasks) > 1:
            mapping["task_id"] = ambiguous_tasks[-1]
    elif "задачи по релизу" in values.values():
        mapping.setdefault("task", mapping.get("epic", 1))
        if ambiguous_tasks:
            mapping["task_id"] = ambiguous_tasks[-1]
    elif "task" in mapping and ambiguous_tasks:
        mapping["task_id"] = ambiguous_tasks[-1]

    score = sum(key in mapping for key in ("epic", "task", "estimate", "executor", "direction"))
    if score < 3 or "estimate" not in mapping:
        return None
    return mapping


def header_mappings(sheet: Any, last_row: int, last_column: int) -> list[tuple[int, dict[str, int]]]:
    matches: list[tuple[int, dict[str, int]]] = []
    for row_number in range(1, last_row + 1):
        mapping = header_mapping(sheet, row_number, last_column)
        if mapping:
            matches.append((row_number, mapping))
    return matches


def cell_payload(formula_cell: Any, value_cell: Any) -> dict[str, Any]:
    payload: dict[str, Any] = {"value": json_value(value_cell.value)}
    if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
        payload["formula"] = formula_cell.value
    if formula_cell.hyperlink:
        payload["hyperlink"] = formula_cell.hyperlink.target or formula_cell.hyperlink.location
    if formula_cell.comment:
        payload["comment"] = {"author": formula_cell.comment.author, "text": formula_cell.comment.text}
    return payload


def is_summary_row(sheet: Any, row_number: int, columns: dict[str, int], last_column: int) -> bool:
    values = [normalize(sheet.cell(row_number, column).value) for column in range(1, last_column + 1)]
    if not any(value in SUMMARY_MARKERS for value in values):
        return False
    estimate = sheet.cell(row_number, columns["estimate"]).value if "estimate" in columns else None
    direction = normalize(sheet.cell(row_number, columns["direction"]).value) if "direction" in columns else ""
    task_id = sheet.cell(row_number, columns["task_id"]).value if "task_id" in columns else None
    has_estimate = isinstance(estimate, (int, float)) and not isinstance(estimate, bool)
    has_direction = direction in {"back", "front", "full", "common", "общая", "бэк", "фронт"}
    has_task_id = bool(task_id and TASK_ID_PATTERN.search(str(task_id)))
    return not (has_estimate or has_direction or has_task_id)


def task_rows(
    formula_sheet: Any,
    value_sheet: Any,
    header_row: int,
    columns: dict[str, int],
    end_row: int,
    last_column: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    tasks: list[dict[str, Any]] = []
    placeholders: list[dict[str, Any]] = []
    block_end = end_row
    for row_number in range(header_row + 1, end_row + 1):
        if is_summary_row(value_sheet, row_number, columns, last_column):
            block_end = row_number - 1
            break
        identity_fields = ("epic", "task", "subtask", "task_id")
        if not any(
            value_sheet.cell(row_number, columns[field]).value not in (None, "")
            for field in identity_fields
            if field in columns
        ):
            continue
        row: dict[str, Any] = {
            "row": row_number,
            "hidden": bool(formula_sheet.row_dimensions[row_number].hidden),
        }
        for field, column in columns.items():
            row[field] = cell_payload(formula_sheet.cell(row_number, column), value_sheet.cell(row_number, column))
        work_fields = ("estimate", "actual_estimate", "direction", "executor", "task_id")
        if not any(row.get(field, {}).get("value") not in (None, "") for field in work_fields):
            row["unplanned"] = True
            placeholders.append(row)
        tasks.append(row)
    return tasks, placeholders, block_end


def is_sprint_sheet(sheet: Any, tasks: list[dict[str, Any]]) -> bool:
    return bool(tasks) and "спринт" in normalize(sheet.title)


def hidden_columns(sheet: Any, last_column: int) -> list[str]:
    result: list[str] = []
    for column in range(1, last_column + 1):
        letter = get_column_letter(column)
        if sheet.column_dimensions[letter].hidden:
            result.append(letter)
    return result


def validations(sheet: Any) -> list[dict[str, Any]]:
    return [
        {
            "ranges": str(validation.sqref),
            "type": validation.type,
            "operator": validation.operator,
            "formula1": json_value(validation.formula1),
            "formula2": json_value(validation.formula2),
            "allow_blank": validation.allow_blank,
            "error": validation.error,
            "prompt": validation.prompt,
        }
        for validation in sheet.data_validations.dataValidation
    ]


def cell_inventory(formula_sheet: Any, value_sheet: Any, last_row: int, last_column: int) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {
        "hyperlinks": [],
        "url_cells": [],
        "formulas": [],
        "formula_errors": [],
        "sumif_range_mismatches": [],
        "comments": [],
    }
    for row in formula_sheet.iter_rows(max_row=last_row or 1, max_col=last_column or 1):
        for cell in row:
            cached = value_sheet[cell.coordinate].value
            if cell.hyperlink:
                result["hyperlinks"].append({
                    "cell": cell.coordinate,
                    "target": cell.hyperlink.target,
                    "location": cell.hyperlink.location,
                    "display": json_value(cell.value),
                })
            extracted_urls = [clean_url(url) for url in URL_PATTERN.findall(str(cell.value))] if cell.value not in (None, "") else []
            if extracted_urls:
                result["url_cells"].append({"cell": cell.coordinate, "urls": extracted_urls})
            if cell.comment:
                result["comments"].append({
                    "cell": cell.coordinate,
                    "author": cell.comment.author,
                    "text": cell.comment.text,
                })
            if not isinstance(cell.value, str) or not cell.value.startswith("="):
                continue
            formula_record = {
                "cell": cell.coordinate,
                "formula": cell.value,
                "cached_value": json_value(cached),
            }
            result["formulas"].append(formula_record)
            cached_error = isinstance(cached, str) and FORMULA_ERROR_PATTERN.search(cached)
            if FORMULA_ERROR_PATTERN.search(cell.value) or cached_error:
                result["formula_errors"].append(formula_record)
            for criteria_range, sum_range in SUMIF_PATTERN.findall(cell.value):
                criteria_rows = range_rows(criteria_range)
                sum_rows = range_rows(sum_range)
                if criteria_rows and sum_rows and criteria_rows != sum_rows:
                    result["sumif_range_mismatches"].append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                        "criteria_rows": list(criteria_rows),
                        "sum_rows": list(sum_rows),
                    })
    return result


def range_rows(reference: str) -> tuple[int, int] | None:
    cleaned = reference.strip().split("!")[-1].replace("$", "")
    try:
        _, min_row, _, max_row = range_boundaries(cleaned)
    except ValueError:
        return None
    return min_row, max_row


def capacity_blocks(formula_sheet: Any, value_sheet: Any, last_row: int, last_column: int) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for header_row in range(1, last_row + 1):
        header_values = [normalize(formula_sheet.cell(header_row, column).value) for column in range(1, last_column + 1)]
        if "часы в спринт" in header_values:
            start_column = header_values.index("часы в спринт") + 1
        elif "разработка" in header_values and any(value.startswith("всего (заложено)") for value in header_values):
            start_column = header_values.index("разработка") + 1
        else:
            continue
        end_column = max(
            column
            for column in range(start_column, last_column + 1)
            if formula_sheet.cell(header_row, column).value not in (None, "")
        )
        end_row = header_row
        empty_streak = 0
        for row_number in range(header_row + 1, min(last_row, header_row + 30) + 1):
            if any(
                formula_sheet.cell(row_number, column).value not in (None, "")
                for column in range(start_column, end_column + 1)
            ):
                end_row = row_number
                empty_streak = 0
            else:
                empty_streak += 1
                if empty_streak >= 2:
                    break
        rows = []
        for row_number in range(header_row, end_row + 1):
            cells = {}
            for column in range(start_column, end_column + 1):
                formula_cell = formula_sheet.cell(row_number, column)
                value_cell = value_sheet.cell(row_number, column)
                if formula_cell.value not in (None, "") or value_cell.value not in (None, ""):
                    cells[formula_cell.coordinate] = cell_payload(formula_cell, value_cell)
            if cells:
                rows.append({"row": row_number, "cells": cells})
        blocks.append({
            "header_row": header_row,
            "end_row": end_row,
            "start_column": start_column,
            "end_column": end_column,
            "rows": rows,
        })
    return blocks


def task_context(sheet: str, task: dict[str, Any]) -> dict[str, Any]:
    return {
        "sheet": sheet,
        "row": task["row"],
        **{
            field: task[field]["value"]
            for field in ("epic", "task", "subtask")
            if field in task
        },
    }


def workbook_audit(path: Path) -> dict[str, Any]:
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    sheets: list[dict[str, Any]] = []
    sprint_count = 0
    task_table_count = 0
    all_task_count = 0
    sprint_task_count = 0
    task_ids: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for index, formula_sheet in enumerate(formulas.worksheets, start=1):
        value_sheet = values[formula_sheet.title]
        last_row, last_column, nonempty_rows = meaningful_bounds(formula_sheet)
        mappings = header_mappings(formula_sheet, last_row, last_column)
        blocks: list[dict[str, Any]] = []
        tasks: list[dict[str, Any]] = []
        for mapping_index, (header_row, columns) in enumerate(mappings):
            next_header = mappings[mapping_index + 1][0] if mapping_index + 1 < len(mappings) else last_row + 1
            rows, placeholders, block_end = task_rows(
                formula_sheet,
                value_sheet,
                header_row,
                columns,
                next_header - 1,
                last_column,
            )
            if rows or placeholders:
                blocks.append({
                    "header_row": header_row,
                    "end_row": block_end,
                    "columns": columns,
                    "tasks": rows,
                    "unplanned_rows": placeholders,
                })
                tasks.extend(rows)
        is_task_table = bool(tasks)
        is_sprint = is_sprint_sheet(formula_sheet, tasks)
        sprint_count += int(is_sprint)
        task_table_count += int(is_task_table)
        all_task_count += len(tasks)
        sprint_task_count += len(tasks) if is_sprint else 0
        for task in tasks:
            task_id = task.get("task_id", {}).get("value")
            if task_id:
                match = TASK_ID_PATTERN.search(str(task_id))
                if match:
                    task_ids[match.group(0).upper()].append(task_context(formula_sheet.title, task))
        inventory = cell_inventory(formula_sheet, value_sheet, last_row, last_column)
        sheet: dict[str, Any] = {
            "index": index,
            "title": formula_sheet.title,
            "state": formula_sheet.sheet_state,
            "last_meaningful_row": last_row,
            "last_meaningful_column": last_column,
            "nonempty_rows": nonempty_rows,
            "content_fingerprint": sheet_fingerprint(formula_sheet, last_row, last_column),
            "hidden_rows": [row for row in range(1, last_row + 1) if formula_sheet.row_dimensions[row].hidden],
            "hidden_columns": hidden_columns(formula_sheet, last_column),
            "data_validations": validations(formula_sheet),
            "capacity_blocks": capacity_blocks(formula_sheet, value_sheet, last_row, last_column),
            **inventory,
            "hyperlink_cells": len(inventory["hyperlinks"]),
            "formula_cells": len(inventory["formulas"]),
            "is_task_table": is_task_table,
            "is_sprint": is_sprint,
        }
        if is_task_table:
            sheet.update({"task_blocks": blocks, "tasks": tasks})
        sheets.append(sheet)
    return {
        "workbook": str(path.resolve()),
        "sheet_count": len(sheets),
        "task_table_count": task_table_count,
        "sprint_sheet_count": sprint_count,
        "task_row_count": sprint_task_count,
        "all_task_row_count": all_task_count,
        "duplicate_task_ids": {
            task_id: contexts
            for task_id, contexts in sorted(task_ids.items())
            if len(contexts) > 1
        },
        "sheets": sheets,
    }


def compare_history_snapshot(
    audit: dict[str, Any],
    snapshot_path: Path,
    workbook_id: str,
    target_sheet: str | None,
) -> dict[str, Any]:
    errors: list[str] = []
    try:
        snapshot = json.loads(snapshot_path.read_text())
    except (OSError, json.JSONDecodeError) as exc:
        return {"compatible": False, "errors": [str(exc)]}

    if snapshot.get("schema_version") != 2:
        errors.append("schema_version must be 2")
    if snapshot.get("workbook_id") != workbook_id:
        errors.append("workbook_id does not match")
    gates = snapshot.get("quality", {}).get("gates")
    if not isinstance(gates, dict) or not gates or not all(gates.values()):
        errors.append("stored quality gates are missing or false")

    tasks = snapshot.get("task_history")
    record_ids: set[str] = set()
    if not isinstance(tasks, list):
        errors.append("task_history must be a list")
    else:
        task_ids = [task.get("record_id") for task in tasks if isinstance(task, dict)]
        record_ids = {record_id for record_id in task_ids if isinstance(record_id, str)}
        if len(record_ids) != len(tasks):
            errors.append("task_history record_id values must be unique strings")

    for index_name in ("task_index", "title_index", "epic_index"):
        index = snapshot.get(index_name)
        if not isinstance(index, dict):
            errors.append(f"{index_name} must be an object")
            continue
        unresolved = [
            record_id
            for entry in index.values()
            if isinstance(entry, dict)
            for record_id in entry.get("record_ids", [])
            if record_id not in record_ids
        ]
        if unresolved:
            errors.append(f"{index_name} contains unresolved record IDs")

    task_id_pattern = snapshot.get("task_id_pattern")
    try:
        compiled_task_id = re.compile(task_id_pattern) if isinstance(task_id_pattern, str) else None
    except re.error:
        compiled_task_id = None
    if compiled_task_id is None:
        errors.append("task_id_pattern must be a valid regular expression")
    elif isinstance(snapshot.get("task_index"), dict) and any(
        not compiled_task_id.fullmatch(task_id) for task_id in snapshot["task_index"]
    ):
        errors.append("task_index contains an ID outside task_id_pattern")

    sources = snapshot.get("spec_sources")
    lookup = snapshot.get("source_lookup")
    if not isinstance(sources, dict) or not isinstance(lookup, dict):
        errors.append("spec_sources and source_lookup must be objects")
    elif any(source_key not in sources for source_key in lookup.values()):
        errors.append("source_lookup contains unresolved source keys")

    current = {sheet["title"]: sheet["content_fingerprint"] for sheet in audit["sheets"]}
    stored_sheets = snapshot.get("sheets")
    if not isinstance(stored_sheets, dict):
        errors.append("sheets must be an object")
        stored: dict[str, str | None] = {}
    else:
        stored = {
            title: record.get("content_fingerprint") if isinstance(record, dict) else None
            for title, record in stored_sheets.items()
        }
    compared_titles = (set(current) | set(stored)) - ({target_sheet} if target_sheet else set())
    matched = sorted(title for title in compared_titles if title in current and current[title] == stored.get(title))
    changed = sorted(title for title in compared_titles if title in current and title in stored and current[title] != stored[title])
    new = sorted(title for title in compared_titles if title in current and title not in stored)
    missing = sorted(title for title in compared_titles if title in stored and title not in current)
    return {
        "compatible": not errors,
        "errors": errors,
        "target_sheet_excluded": target_sheet,
        "matched_sheets": matched,
        "changed_sheets": changed,
        "new_sheets": new,
        "missing_sheets": missing,
        "refresh_required": bool(changed or new or missing),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--pretty", action="store_true")
    parser.add_argument("--snapshot", type=Path)
    parser.add_argument("--workbook-id")
    parser.add_argument("--target-sheet")
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"workbook not found: {args.workbook}")
    if args.snapshot and not args.snapshot.is_file():
        parser.error(f"snapshot not found: {args.snapshot}")
    if args.snapshot and not args.workbook_id:
        parser.error("--workbook-id is required with --snapshot")
    try:
        result = workbook_audit(args.workbook)
    except Exception as exc:  # noqa: BLE001 - CLI must report malformed workbooks
        print(f"audit_workbook.py: {exc}", file=sys.stderr)
        return 1
    if args.snapshot:
        result["snapshot_comparison"] = compare_history_snapshot(
            result,
            args.snapshot,
            args.workbook_id,
            args.target_sheet,
        )
    json.dump(result, sys.stdout, ensure_ascii=False, indent=2 if args.pretty else None)
    sys.stdout.write("\n")
    return 2 if args.snapshot and not result["snapshot_comparison"]["compatible"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
